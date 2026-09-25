"""
reports/excel_report.py

Membuat file Excel (.xlsx) berisi riwayat seluruh hasil evaluasi,
dengan Zone A dan Zone C terpisah.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "reports")


def generate_excel_riwayat(data: list) -> str:
    os.makedirs(REPORTS_DIR, exist_ok=True)
    path = os.path.join(REPORTS_DIR, "riwayat_hasil_ndt.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Riwayat Hasil NDT"

    headers = [
        "ID", "Jenis Benda", "Diameter (mm)", "Zona",
        "Luas Babbit (mm2)", "Luas Zone A (mm2)", "Luas Zone C (mm2)",
        "Unbond Zone A (mm2)", "% Unbond Zone A", "Status Zone A",
        "Unbond Zone C (mm2)", "% Unbond Zone C", "Status Zone C",
        "Status Akhir",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in data:
        ws.append([
            row["id"], row["jenis_benda"], row["diameter_mm"], row["zona"],
            row["a_babbit"], row["a_zone_a"], row["a_zone_c"],
            row["a_unbond_zone_a"], row["persen_unbond_zone_a"], row["status_zone_a"],
            row["a_unbond_zone_c"], row["persen_unbond_zone_c"], row["status_zone_c"],
            row["status"],
        ])

        for col_idx in (10, 13, 14):  # Status Zone A, Status Zone C, Status Akhir
            cell = ws.cell(row=ws.max_row, column=col_idx)
            if cell.value == "ACCEPT":
                cell.fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
            else:
                cell.fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    wb.save(path)
    return path